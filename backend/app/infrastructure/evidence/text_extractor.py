"""
File: app/infrastructure/evidence/text_extractor.py
Folder: backend/app/infrastructure/evidence
Purpose: Turn an uploaded evidence file into plain text the LLM can read.
         Supports the formats TISAX evidence actually comes in:
           - PDF  (pypdf; flags image-only/scanned PDFs)
           - DOCX (python-docx, including tables)
           - TXT  (decoded directly)
           - images PNG/JPG/TIFF (OCR via tesseract)

         When a PDF has no text layer (a scanned/signed policy), we can't
         read it here, so we flag it clearly and tell the user to upload a
         text PDF or an image (which goes through OCR) instead.
"""

import io
import logging

from app.core.exceptions import ValidationError

logger = logging.getLogger(__name__)

SUPPORTED_MIME = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "image/png": "image",
    "image/jpeg": "image",
    "image/tiff": "image",
}

MAX_TEXT_CHARS = 20000  # cap what we forward to the LLM


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    pages_text = []
    empty_pages = 0
    for page in reader.pages:
        txt = page.extract_text() or ""
        if not txt.strip():
            empty_pages += 1
        pages_text.append(txt)
    text = "\n".join(pages_text).strip()
    if not text and empty_pages:
        raise ValidationError(
            "This PDF appears to be scanned (no text layer). Upload a text-based "
            "PDF, or export the page as an image so it can be processed by OCR."
        )
    return text


def _extract_docx(content: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _extract_image(content: bytes) -> str:
    import pytesseract
    from PIL import Image

    from app.core.config import get_settings

    tesseract_cmd = getattr(get_settings(), "TESSERACT_CMD", "")
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    try:
        image = Image.open(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001
        raise ValidationError("Could not open the uploaded image.") from exc
    text = pytesseract.image_to_string(image)
    return text.strip()


def extract_text(content: bytes, mime_type: str) -> str:
    kind = SUPPORTED_MIME.get(mime_type)
    if kind is None:
        raise ValidationError(
            f"Unsupported file type '{mime_type}'. Accepted: PDF, DOCX, TXT, PNG, JPG, TIFF."
        )

    if kind == "pdf":
        text = _extract_pdf(content)
    elif kind == "docx":
        text = _extract_docx(content)
    elif kind == "txt":
        text = content.decode("utf-8", errors="replace").strip()
    else:  # image
        text = _extract_image(content)

    if not text:
        raise ValidationError("No readable text could be extracted from the file.")

    if len(text) > MAX_TEXT_CHARS:
        logger.info("Evidence text truncated from %d to %d chars", len(text), MAX_TEXT_CHARS)
        text = text[:MAX_TEXT_CHARS] + "\n[...truncated...]"
    return text